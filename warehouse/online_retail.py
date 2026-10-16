"""The real-data feed: UCI Online Retail II, split into one CSV per day in the landing folder.

The dataset is a real UK online gift shop's invoice lines from 2009-12-01 to 2011-12-09, about
1.07 million rows (Chen, D. 2012, UCI Machine Learning Repository, CC BY 4.0,
https://archive.ics.uci.edu/dataset/502/online+retail+ii). It is downloaded, never committed.

    python -m warehouse.online_retail                         # download (about 45 MB) and land every day
    python -m warehouse.online_retail --start 2010-12-01 --days 14   # land two weeks only
    python -m warehouse.online_retail --file online_retail_II.xlsx   # use a file you downloaded yourself

The workbook has one sheet per year and the sheets can overlap at the year boundary; when two sheets
both have a day, the one with more lines for that day is kept, so no day is counted twice.
After landing, `python -m warehouse.run` checks and loads the files like any other.
"""
import argparse
import csv
import io
import os
import shutil
import urllib.request
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta

from warehouse import config, feeds

URL = os.environ.get("ONLINE_RETAIL_URL", "https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip")
DOWNLOAD_DIR = os.path.join(config.DATA_DIR, "downloads")
FEED = feeds.ONLINE_RETAIL
# the workbook's headers, in the order of FEED.columns
SOURCE_HEADERS = ["invoice", "stockcode", "description", "quantity", "invoicedate", "price", "customer id", "country"]


def download(url=URL, folder=DOWNLOAD_DIR):
    """Downloads the zip once and keeps it. Returns its path."""
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, "online_retail_ii.zip")
    if not os.path.exists(path):
        tmp = path + ".part"
        req = urllib.request.Request(url, headers={"User-Agent": "retail-warehouse"})
        with urllib.request.urlopen(req, timeout=300) as r, open(tmp, "wb") as f:  # nosec - fixed public URL
            shutil.copyfileobj(r, f)
        os.replace(tmp, path)
    return path


def workbook_bytes(path):
    """The .xlsx itself, or the .xlsx inside the downloaded zip (an .xlsx is a zip too, so look inside)."""
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            inner = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if inner:
                return z.read(inner[0])
    with open(path, "rb") as f:
        return f.read()


def text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_sheets(path):
    """Yields (sheet name, {day: [row, ...]}) for each sheet, rows in FEED.columns order as text."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(workbook_bytes(path)), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = [text(h).lower() for h in next(rows, [])]
            missing = [h for h in SOURCE_HEADERS if h not in header]
            if missing:
                raise ValueError(f"Sheet {ws.title} is missing columns: {', '.join(missing)}")
            idx = [header.index(h) for h in SOURCE_HEADERS]
            by_day = defaultdict(list)
            for row in rows:
                when = row[idx[4]]
                if not isinstance(when, datetime):
                    continue                  # blank trailing rows
                by_day[when.date()].append([text(row[i]) for i in idx])
            yield ws.title, by_day
    finally:
        wb.close()


def pick_days(sheets):
    """One set of rows per day; where sheets overlap, the sheet with more rows for that day wins."""
    chosen = {}
    for _name, by_day in sheets:
        for day, rows in by_day.items():
            if day not in chosen or len(rows) > len(chosen[day]):
                chosen[day] = rows
    return chosen


def write_days(days, out_dir, start=None, count=None):
    """Writes online_retail_<day>.csv for each chosen day in range. Returns (files, rows)."""
    os.makedirs(out_dir, exist_ok=True)
    keep = sorted(days)
    if start:
        keep = [d for d in keep if d >= start]
    if count:
        keep = [d for d in keep if d < (start or keep[0]) + timedelta(days=count)] if keep else keep
    total = 0
    for day in keep:
        with open(os.path.join(out_dir, FEED.file_name(day)), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(FEED.columns)
            w.writerows(days[day])
        total += len(days[day])
    return len(keep), total


def land(out_dir=None, path=None, start=None, count=None):
    path = path or download()
    return write_days(pick_days(read_sheets(path)), out_dir or config.LANDING_DIR, start, count)


def landed_days(out_dir):
    names = os.listdir(out_dir) if os.path.isdir(out_dir) else []
    return sorted(FEED.day_of(n) for n in names if FEED.owns(n))


def main(argv=None):
    p = argparse.ArgumentParser(description="Land the UCI Online Retail II dataset as daily files.")
    p.add_argument("--file", help="a downloaded online_retail_II.xlsx (or the zip); default: download it")
    p.add_argument("--out", default=config.LANDING_DIR, help="folder for the daily files")
    p.add_argument("--start", type=date.fromisoformat, help="first day to land (YYYY-MM-DD)")
    p.add_argument("--days", type=int, help="how many calendar days to land from --start")
    a = p.parse_args(argv)
    files, rows = land(a.out, a.file, a.start, a.days)
    print(f"Wrote {files:,} daily files with {rows:,} invoice lines to {a.out}")


if __name__ == "__main__":
    main()
